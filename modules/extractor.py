# modules/extractor.py
import pdfplumber
from docx import Document

def extract_text(file, file_type):
    if file_type == "pdf":
        return extract_from_pdf(file)
    elif file_type == "docx":
        return extract_from_docx(file)
    elif file_type == "doc":
        return extract_from_docx(file)
    elif file_type == "txt":
        return file.read().decode("utf-8", errors="ignore")
    else:
        return ""

def extract_from_pdf(file):
    text = ""
    with pdfplumber.open(file) as pdf:
        all_text = []
        for page in pdf.pages:
            words = page.extract_words(x_tolerance=2, y_tolerance=3)
            if not words:
                all_text.append(page.extract_text() or "")
                continue

            pw = page.width
            mid_start = int(pw * 0.20)
            mid_end = int(pw * 0.80)
            word_boxes = [(w['x0'], w['x1'], w['top'], w['bottom']) for w in words]

            best_split = None
            best_gutter_width = 0

            for x in range(mid_start, mid_end, 5):
                intersecting = [b for b in word_boxes if not (b[1] < x - 5 or b[0] > x + 5)]
                if len(intersecting) == 0:
                    left_edge = max([b[1] for b in word_boxes if b[1] <= x], default=0)
                    right_edge = min([b[0] for b in word_boxes if b[0] >= x], default=pw)
                    gutter = right_edge - left_edge
                    if gutter >= 10 and gutter > best_gutter_width:
                        left_cnt = sum(1 for b in word_boxes if b[1] <= x)
                        right_cnt = sum(1 for b in word_boxes if b[0] >= x)
                        if left_cnt >= len(words) * 0.15 and right_cnt >= len(words) * 0.15:
                            best_gutter_width = gutter
                            best_split = (left_edge + right_edge) / 2

            def group_lines(w_list):
                if not w_list:
                    return ""
                w_list.sort(key=lambda w: (w['top'], w['x0']))
                lines = []
                curr_line = [w_list[0]]
                for w in w_list[1:]:
                    if abs(w['top'] - curr_line[-1]['top']) <= 4:
                        curr_line.append(w)
                    else:
                        curr_line.sort(key=lambda x: x['x0'])
                        lines.append(" ".join(x['text'] for x in curr_line))
                        curr_line = [w]
                if curr_line:
                    curr_line.sort(key=lambda x: x['x0'])
                    lines.append(" ".join(x['text'] for x in curr_line))
                return "\n".join(lines)

            if best_split is not None:
                left_words = [w for w in words if w['x1'] <= best_split + 5]
                right_words = [w for w in words if w['x0'] >= best_split - 5]
                middle_words = [w for w in words if w['x0'] < best_split - 5 and w['x1'] > best_split + 5]

                page_lines = []
                top_span = [w for w in middle_words if w['top'] < min(left_words[0]['top'] if left_words else 9999, right_words[0]['top'] if right_words else 9999)]
                if top_span:
                    page_lines.append(group_lines(top_span))
                if left_words:
                    page_lines.append(group_lines(left_words))
                if right_words:
                    page_lines.append(group_lines(right_words))
                other_middle = [w for w in middle_words if w not in top_span]
                if other_middle:
                    page_lines.append(group_lines(other_middle))
                all_text.append("\n".join(page_lines))
            else:
                all_text.append(group_lines(words))
        return "\n\n".join(all_text)

def extract_from_docx(file):
    doc = Document(file)
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += cell.text + " "
    return text